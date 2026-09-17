from openpyxl import load_workbook
wb = load_workbook("inventory_sample.xlsx")   # 워크북(파일 전체)
print(wb.sheetnames)                          # 시트 이름 목록
ws = wb["재고현황"]                            # 워크시트 선택
print(ws["B2"].value)                         # 셀 단위 접근
for row in ws.iter_rows(min_row=2, values_only=True):
    print(row)                                # 튜플로 행 순회
